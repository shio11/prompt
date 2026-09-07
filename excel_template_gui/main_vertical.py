"""Excelテンプレート入力GUIアプリ（縦バージョン／複数レコード対応）

template_vertical.xlsx のように、
    A列: 項目名
    B列以降: レコードごとの値（値1, 値2, 値3, ... 列数は可変）
という「縦（項目が行方向に並ぶ）」レイアウトのExcelを対象にしたバージョンです。

1つのファイルに複数人（複数レコード）分のデータが列として並んでいる場合に対応し、
GUI上では1レコードずつフォーム表示して「前へ／次へ」で切り替えられます。
"""

from __future__ import annotations

import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FIELDS: Final[list[str]] = [
    "氏名",
    "フリガナ",
    "会社名",
    "部署",
    "メールアドレス",
    "電話番号",
    "日付",
    "備考",
]

TEMPLATE_PATH: Final[Path] = Path(__file__).parent / "template_vertical.xlsx"

# 縦レイアウトのヘッダー（1行目）
ITEM_HEADER: Final[str] = "項目"
VALUE_HEADER_PREFIX: Final[str] = "値"

Record = dict[str, str]


def ensure_template(path: Path) -> None:
    """テンプレートファイルが存在しなければ、縦レイアウトの新規Excelを作成する

    1行目: 項目 / 値1 のヘッダー
    2行目以降: FIELDS の各項目名（A列）と、空欄の値（B列）
    最初は1レコード分のみの空テンプレートとして作成する。
    """
    if path.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "template"

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet["A1"] = ITEM_HEADER
    sheet["B1"] = f"{VALUE_HEADER_PREFIX}1"
    for cell in (sheet["A1"], sheet["B1"]):
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row_index, field in enumerate(FIELDS, start=2):
        label_cell = sheet.cell(row=row_index, column=1, value=field)
        value_cell = sheet.cell(row=row_index, column=2, value="")
        label_cell.font = header_font
        label_cell.fill = header_fill
        label_cell.border = border
        value_cell.border = border

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 30

    workbook.save(path)


def load_template_records(path: Path) -> list[Record]:
    """テンプレートを読み込み、レコード（列）ごとの辞書のリストを返す"""
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    return _read_labeled_records(sheet)


def read_external_file(path: Path) -> list[Record]:
    """外部Excelファイルを読み込み、レコード（列）ごとの辞書のリストを返す

    A列に項目名、B列以降に値が並ぶ「縦レイアウト・複数列」を主に想定しているが、
    1行目が見出し・2行目のみがデータという「横レイアウト・単一レコード」の
    ファイルが渡された場合も、1件のレコードとして読み込めるようにフォールバックする。
    """
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    return _read_labeled_records(sheet)


def write_output_file(path: Path, records: list[Record]) -> None:
    """複数レコードを縦レイアウトのExcelとして書き出す

    A列: 項目名（FIELDS の並び順）
    B列以降: 各レコードの値（値1, 値2, ... のヘッダー付き）
    """
    if not records:
        records = [{}]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "output"

    sheet["A1"] = ITEM_HEADER
    for record_index in range(len(records)):
        sheet.cell(row=1, column=2 + record_index, value=f"{VALUE_HEADER_PREFIX}{record_index + 1}")

    for row_index, field in enumerate(FIELDS, start=2):
        sheet.cell(row=row_index, column=1, value=field)
        for record_index, record in enumerate(records):
            sheet.cell(
                row=row_index,
                column=2 + record_index,
                value=record.get(field, ""),
            )

    sheet.column_dimensions["A"].width = 16
    for record_index in range(len(records)):
        col_letter = sheet.cell(row=1, column=2 + record_index).column_letter
        sheet.column_dimensions[col_letter].width = 20

    workbook.save(path)


def _is_usable_value(raw_value: object) -> bool:
    """セルの値が「データ」として採用できるかを判定する

    - 空セル（None）や空文字列は不採用
    - 値がたまたま FIELDS の項目名（例: 別の見出しセル）と一致する場合も不採用
      （見出しセル同士が隣り合っているケースで、見出しを値として誤読するのを防ぐ）
    """
    if raw_value is None:
        return False
    text = str(raw_value).strip()
    return text != "" and text not in FIELDS


def _read_labeled_records(sheet: Worksheet) -> list[Record]:
    """シートを走査し、FIELDSと一致するセルを見つけて複数レコードを読み取る

    優先: 縦レイアウト・複数列（列位置は問わない）
        各行を左から走査し、FIELDSのいずれかと一致する最初の未使用セルを
        「項目名セル」とみなす。項目名セルが何列目にあっても構わない
        （例: A列開始でもD列開始でも対応）。項目名セルより右側にある
        セルを「レコード1, レコード2, ...」の値として順番に採用する
        （列数は可変）。値の末尾に続く空セルは切り詰めるが、途中の空セルは
        空文字として保持し、レコードの列位置がズレないようにする。

    フォールバック: 横レイアウト・単一レコード
        縦レイアウトで1件も値が見つからなかった場合、1行目を見出し・2行目を
        データ行とみなす従来方式で1件だけ読み込む。
    """
    field_row_values: dict[str, list[str]] = {}
    max_len = 0

    for row in sheet.iter_rows():
        if not row:
            continue

        label_cell = None
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            if text in FIELDS and text not in field_row_values:
                label_cell = cell
                break

        if label_cell is None:
            continue

        label = str(label_cell.value).strip()
        label_row, label_col = label_cell.row, label_cell.column

        raw_values = [
            sheet.cell(row=label_row, column=col).value
            for col in range(label_col + 1, sheet.max_column + 1)
        ]
        while raw_values and raw_values[-1] is None:
            raw_values.pop()

        values = [str(v) if _is_usable_value(v) else "" for v in raw_values]
        field_row_values[label] = values
        max_len = max(max_len, len(values))

    has_any_value = any(v.strip() for values in field_row_values.values() for v in values)

    if field_row_values and has_any_value:
        records: list[Record] = []
        for i in range(max_len):
            record: Record = {}
            for field in FIELDS:
                values = field_row_values.get(field, [])
                record[field] = values[i] if i < len(values) else ""
            records.append(record)
        return records

    # フォールバック: 横レイアウト（1行目=見出し、2行目=データ）の単一レコード
    header = [str(cell.value) if cell.value is not None else "" for cell in sheet[1]]
    fallback: Record = {}
    if sheet.max_row >= 2:
        for col_index, name in enumerate(header, start=1):
            if name in FIELDS:
                cell_value = sheet.cell(row=2, column=col_index).value
                fallback[name] = "" if cell_value is None else str(cell_value)
    return [fallback] if fallback else [{}]


class TemplateApp:
    """テンプレート入力用のGUIアプリケーション（縦レイアウト・複数レコード対応）"""

    def __init__(self, root: tk.Tk) -> None:
        self._root = root
        self._entries: dict[str, tk.Entry] = {}
        self._records: list[Record] = [{}]
        self._current_index: int = 0
        self._position_label: tk.Label | None = None
        self._build_widgets()
        self._load_initial_values()

    def _build_widgets(self) -> None:
        self._root.title("Excelテンプレート入力（縦バージョン・複数レコード対応）")

        form_frame = tk.Frame(self._root, padx=16, pady=16)
        form_frame.pack(fill="both", expand=True)

        for row, field in enumerate(FIELDS):
            label = tk.Label(form_frame, text=field, width=12, anchor="w")
            label.grid(row=row, column=0, sticky="w", pady=4)

            entry = tk.Entry(form_frame, width=40)
            entry.grid(row=row, column=1, pady=4, padx=(8, 0))
            self._entries[field] = entry

        # レコード切り替えナビゲーション
        nav_frame = tk.Frame(self._root, padx=16, pady=4)
        nav_frame.pack(fill="x")

        prev_button = tk.Button(nav_frame, text="◀ 前へ", command=self._on_prev)
        prev_button.pack(side="left")

        self._position_label = tk.Label(nav_frame, text="0 / 0")
        self._position_label.pack(side="left", padx=12)

        next_button = tk.Button(nav_frame, text="次へ ▶", command=self._on_next)
        next_button.pack(side="left")

        add_button = tk.Button(nav_frame, text="＋ 新規レコード追加", command=self._on_add_record)
        add_button.pack(side="left", padx=(16, 0))

        button_frame = tk.Frame(self._root, padx=16, pady=8)
        button_frame.pack(fill="x")

        import_button = tk.Button(
            button_frame, text="外部ファイルを読み込む", command=self._on_import
        )
        import_button.pack(side="left", padx=(0, 8))

        export_button = tk.Button(button_frame, text="出力する", command=self._on_export)
        export_button.pack(side="left")

    def _load_initial_values(self) -> None:
        try:
            records = load_template_records(TEMPLATE_PATH)
        except Exception as error:
            messagebox.showerror("エラー", f"テンプレートの読み込みに失敗しました: {error}")
            return
        self._records = records if records else [{}]
        self._current_index = 0
        self._refresh_form()

    def _refresh_form(self) -> None:
        """現在のインデックスのレコードをフォームに反映し、位置表示を更新する"""
        total = len(self._records)
        if total == 0:
            self._records = [{}]
            total = 1
        self._current_index = max(0, min(self._current_index, total - 1))

        self._set_values(self._records[self._current_index])
        if self._position_label is not None:
            self._position_label.config(text=f"{self._current_index + 1} / {total}")

    def _set_values(self, values: Record) -> None:
        for field, entry in self._entries.items():
            entry.delete(0, tk.END)
            entry.insert(0, values.get(field, ""))

    def _get_values(self) -> Record:
        return {field: entry.get() for field, entry in self._entries.items()}

    def _save_current_entry_to_records(self) -> None:
        """フォームの編集内容を、現在のインデックスのレコードに書き戻す"""
        if not self._records:
            self._records = [{}]
        self._records[self._current_index] = self._get_values()

    def _on_prev(self) -> None:
        if self._current_index <= 0:
            return
        self._save_current_entry_to_records()
        self._current_index -= 1
        self._refresh_form()

    def _on_next(self) -> None:
        if self._current_index >= len(self._records) - 1:
            return
        self._save_current_entry_to_records()
        self._current_index += 1
        self._refresh_form()

    def _on_add_record(self) -> None:
        self._save_current_entry_to_records()
        self._records.append({})
        self._current_index = len(self._records) - 1
        self._refresh_form()

    def _on_import(self) -> None:
        file_path = filedialog.askopenfilename(
            title="読み込む外部ファイルを選択",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not file_path:
            return

        try:
            records = read_external_file(Path(file_path))
        except Exception as error:
            messagebox.showerror("エラー", f"ファイルの読み込みに失敗しました: {error}")
            return

        if not records or not any(any(v for v in record.values()) for record in records):
            messagebox.showwarning(
                "警告", "テンプレートの項目と一致するデータが見つかりませんでした"
            )
            return

        self._records = records
        self._current_index = 0
        self._refresh_form()
        messagebox.showinfo(
            "完了", f"外部ファイルから {len(records)} 件のレコードを読み込みました"
        )

    def _on_export(self) -> None:
        file_path = filedialog.asksaveasfilename(
            title="出力先ファイルを指定",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not file_path:
            return

        self._save_current_entry_to_records()

        try:
            write_output_file(Path(file_path), self._records)
        except Exception as error:
            messagebox.showerror("エラー", f"出力に失敗しました: {error}")
            return

        messagebox.showinfo("完了", f"{len(self._records)} 件のレコードを出力しました: {file_path}")


def main() -> None:
    ensure_template(TEMPLATE_PATH)
    root = tk.Tk()
    TemplateApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
