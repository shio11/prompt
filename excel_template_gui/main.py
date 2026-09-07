"""Excelテンプレート入力GUIアプリ"""

from __future__ import annotations

import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

FIELDS: Final[list[str]] = [
    "氏名",
    "フリガナ",
    "会社名",
    "部署",
    "メールアドレス",
    "電話番号",
    "日付",
    "備考",
]

TEMPLATE_PATH: Final[Path] = Path(__file__).parent / "template.xlsx"


def ensure_template(path: Path) -> None:
    """テンプレートファイルが存在しなければ、ヘッダー付きの新規Excelを作成する"""
    if path.exists():
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "template"
    sheet.append(FIELDS)
    workbook.save(path)


def load_template_values(path: Path) -> dict[str, str]:
    """テンプレートの1行目のデータをフィールド名でマッピングして返す"""
    workbook = load_workbook(path)
    sheet = workbook.active
    return _read_first_data_row(sheet)


def read_external_file(path: Path) -> dict[str, str]:
    """外部Excelファイルのヘッダーとデータ行を読み込み、FIELDSに一致する値を返す"""
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    return _read_first_data_row(sheet)


def write_output_file(path: Path, values: dict[str, str]) -> None:
    """入力値をテンプレート形式のExcelとして書き出す"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "output"
    sheet.append(FIELDS)
    sheet.append([values.get(field, "") for field in FIELDS])
    workbook.save(path)


def _read_first_data_row(sheet: Worksheet) -> dict[str, str]:
    header = [str(cell.value) if cell.value is not None else "" for cell in sheet[1]]
    values: dict[str, str] = {}
    if sheet.max_row < 2:
        return values
    for col_index, name in enumerate(header, start=1):
        if name in FIELDS:
            cell_value = sheet.cell(row=2, column=col_index).value
            values[name] = "" if cell_value is None else str(cell_value)
    return values


class TemplateApp:
    """テンプレート入力用のGUIアプリケーション"""

    def __init__(self, root: tk.Tk) -> None:
        self._root = root
        self._entries: dict[str, tk.Entry] = {}
        self._build_widgets()
        self._load_initial_values()

    def _build_widgets(self) -> None:
        self._root.title("Excelテンプレート入力")

        form_frame = tk.Frame(self._root, padx=16, pady=16)
        form_frame.pack(fill="both", expand=True)

        for row, field in enumerate(FIELDS):
            label = tk.Label(form_frame, text=field, width=12, anchor="w")
            label.grid(row=row, column=0, sticky="w", pady=4)

            entry = tk.Entry(form_frame, width=40)
            entry.grid(row=row, column=1, pady=4, padx=(8, 0))
            self._entries[field] = entry

        button_frame = tk.Frame(self._root, padx=16, pady=8)
        button_frame.pack(fill="x")

        import_button = tk.Button(
            button_frame, text="外部ファイルを読み込む", command=self._on_import
        )
        import_button.pack(side="left", padx=(0, 8))

        export_button = tk.Button(button_frame, text="出力する", command=self._on_export)
        export_button.pack(side="left")

    def _load_initial_values(self) -> None:
        try:
            values = load_template_values(TEMPLATE_PATH)
        except Exception as error:
            messagebox.showerror("エラー", f"テンプレートの読み込みに失敗しました: {error}")
            return
        self._set_values(values)

    def _set_values(self, values: dict[str, str]) -> None:
        for field, entry in self._entries.items():
            entry.delete(0, tk.END)
            entry.insert(0, values.get(field, ""))

    def _get_values(self) -> dict[str, str]:
        return {field: entry.get() for field, entry in self._entries.items()}

    def _on_import(self) -> None:
        file_path = filedialog.askopenfilename(
            title="読み込む外部ファイルを選択",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not file_path:
            return

        try:
            values = read_external_file(Path(file_path))
        except Exception as error:
            messagebox.showerror("エラー", f"ファイルの読み込みに失敗しました: {error}")
            return

        if not values:
            messagebox.showwarning(
                "警告", "テンプレートの項目と一致する列が見つかりませんでした"
            )
            return

        self._set_values(values)
        messagebox.showinfo("完了", "外部ファイルの内容をテンプレートに反映しました")

    def _on_export(self) -> None:
        file_path = filedialog.asksaveasfilename(
            title="出力先ファイルを指定",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not file_path:
            return

        try:
            write_output_file(Path(file_path), self._get_values())
        except Exception as error:
            messagebox.showerror("エラー", f"出力に失敗しました: {error}")
            return

        messagebox.showinfo("完了", f"出力しました: {file_path}")


def main() -> None:
    ensure_template(TEMPLATE_PATH)
    root = tk.Tk()
    TemplateApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
