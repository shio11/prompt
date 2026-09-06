# -*- coding: utf-8 -*-
"""対象資料の入力元サンプル（別ファイル）を作成する。

支店ごとの経費明細を持つ3シート構成。
Excel転記アプリはこのファイルの各シートから、GUIでチェックされた
支店のデータのみを読み込む想定。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUTPUT_PATH = "source_data.xlsx"

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
BASE_FONT = Font(name="Arial")

# シート名: (表示用資料名, 明細データ[日付, 費目, 金額])
SHEETS = {
    "資料A_東京支店": (
        "東京支店",
        [
            ("2026-08-03", "交通費", 12400),
            ("2026-08-10", "消耗品費", 8600),
            ("2026-08-18", "会議費", 15200),
            ("2026-08-25", "通信費", 5400),
        ],
    ),
    "資料B_大阪支店": (
        "大阪支店",
        [
            ("2026-08-05", "交通費", 9800),
            ("2026-08-12", "接待交際費", 22000),
            ("2026-08-21", "消耗品費", 6300),
        ],
    ),
    "資料C_名古屋支店": (
        "名古屋支店",
        [
            ("2026-08-07", "通信費", 4800),
            ("2026-08-14", "交通費", 11200),
            ("2026-08-20", "会議費", 9600),
            ("2026-08-28", "消耗品費", 7100),
        ],
    ),
}

HEADERS = ["日付", "費目", "金額"]


def build() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, (display_name, rows) in SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = f"{display_name} 経費明細（2026年8月）"
        ws["A1"].font = Font(name="Arial", bold=True, size=12)
        ws.merge_cells("A1:C1")

        for col, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

        for row_offset, (date_str, category, amount) in enumerate(rows, start=4):
            ws.cell(row=row_offset, column=1, value=date_str).font = BASE_FONT
            ws.cell(row=row_offset, column=2, value=category).font = BASE_FONT
            amount_cell = ws.cell(row=row_offset, column=3, value=amount)
            amount_cell.font = BASE_FONT
            amount_cell.number_format = "#,##0"

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 12

    wb.save(OUTPUT_PATH)
    print(f"作成完了: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
