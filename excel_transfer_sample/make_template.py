# -*- coding: utf-8 -*-
"""Excel転記アプリの転記先テンプレート（複数シート構成）サンプルを作成する。

シート構成:
  - 設定  : 対象資料の一覧（GUIのチェックボックス初期状態、入力元ファイル/シート名）
  - サマリー: 作成日・作成者、資料ごとの小計（数式）、総合計（数式）
  - 明細  : GUIが読み込んだ明細データの転記先（Pythonが書き込む対象）

「設定」シートの内容を書き換えれば、対象資料の追加・入力元ファイルの
変更にも対応できる構成にしている。
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUTPUT_PATH = "template.xlsx"

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
INPUT_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
LABEL_FONT = Font(name="Arial", bold=True)
BASE_FONT = Font(name="Arial")
SAMPLE_FONT = Font(name="Arial", italic=True, color="808080")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 資料キー, 表示名, 読込対象の初期値, 入力元ファイル名, 入力元シート名
SOURCE_ROWS = [
    ("資料A", "東京支店", True, "source_data.xlsx", "資料A_東京支店"),
    ("資料B", "大阪支店", True, "source_data.xlsx", "資料B_大阪支店"),
    ("資料C", "名古屋支店", False, "source_data.xlsx", "資料C_名古屋支店"),
]


def build_settings_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("設定")
    ws["A1"] = "対象資料 設定一覧（GUIのチェックボックス初期状態として読み込まれます）"
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws.merge_cells("A1:E1")

    headers = ["資料キー", "資料名（表示用）", "読込対象", "入力元ファイル名", "入力元シート名"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    for row_offset, (key, name, checked, src_file, src_sheet) in enumerate(SOURCE_ROWS, start=4):
        values = [key, name, checked, src_file, src_sheet]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_offset, column=col, value=value)
            cell.font = BASE_FONT
            cell.border = BORDER
            if col == 3:
                cell.fill = INPUT_FILL
                cell.alignment = Alignment(horizontal="center")

    ws["C4"].comment = Comment(
        "TRUE = GUI起動時に読込対象としてチェックON / FALSE = チェックOFF", "template"
    )

    widths = {"A": 10, "B": 16, "C": 10, "D": 20, "E": 20}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_summary_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("サマリー")
    ws["A1"] = "部門別経費 月次サマリー"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws.merge_cells("A1:D1")

    ws["A3"] = "作成日"
    ws["A3"].font = LABEL_FONT
    ws["B3"].fill = INPUT_FILL
    ws["B3"].comment = Comment("Pythonが出力時にセットする欄（記入例）", "template")
    ws["B3"] = "2026-09-07"

    ws["A4"] = "作成者"
    ws["A4"].font = LABEL_FONT
    ws["B4"].fill = INPUT_FILL
    ws["B4"].comment = Comment("Pythonが出力時にセットする欄（記入例）", "template")
    ws["B4"] = "山田太郎"

    headers = ["資料名", "読込対象", "小計（円）"]
    header_row = 6
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    first_data_row = header_row + 1
    for i, (_key, _name, _checked, _src_file, _src_sheet) in enumerate(SOURCE_ROWS):
        row = first_data_row + i
        setting_row = 4 + i  # 設定シートの対応行
        ws.cell(row=row, column=1, value=f"='設定'!B{setting_row}").font = BASE_FONT
        ws.cell(row=row, column=2, value=f"='設定'!C{setting_row}").font = BASE_FONT
        # 明細シートのA列（資料名）が一致する行のD列（金額）を合計する
        ws.cell(
            row=row, column=3,
            value=f"=SUMIFS(明細!$D:$D,明細!$A:$A,A{row})",
        ).font = BASE_FONT
        for col in range(1, 4):
            ws.cell(row=row, column=col).border = BORDER
        ws.cell(row=row, column=3).number_format = "#,##0"

    total_row = first_data_row + len(SOURCE_ROWS)
    ws.cell(row=total_row, column=2, value="総合計").font = LABEL_FONT
    total_cell = ws.cell(
        row=total_row, column=3,
        value=f"=SUM(C{first_data_row}:C{total_row - 1})",
    )
    total_cell.font = Font(name="Arial", bold=True)
    total_cell.number_format = "#,##0"
    for col in range(1, 4):
        ws.cell(row=total_row, column=col).border = BORDER

    widths = {"A": 14, "B": 10, "C": 14}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_detail_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("明細")
    ws["A1"] = "経費明細（GUIから読み込んだ対象資料の明細がここに転記されます）"
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws.merge_cells("A1:D1")

    headers = ["資料名", "日付", "費目", "金額"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    # 記入例（Pythonの出力処理で上書きされる想定の1行）
    sample = ["東京支店", "2026-08-03", "交通費", 12400]
    for col, value in enumerate(sample, start=1):
        cell = ws.cell(row=4, column=col, value=value)
        cell.font = SAMPLE_FONT
        cell.border = BORDER
        if col == 4:
            cell.number_format = "#,##0"
    ws["A4"].comment = Comment(
        "記入例です。Excel出力時にこの範囲はGUIの表示内容で上書きされます。", "template"
    )

    widths = {"A": 14, "B": 14, "C": 14, "D": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    build_settings_sheet(wb)
    build_summary_sheet(wb)
    build_detail_sheet(wb)
    wb.active = wb.sheetnames.index("サマリー")
    wb.save(OUTPUT_PATH)
    print(f"作成完了: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
